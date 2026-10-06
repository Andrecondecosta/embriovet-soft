"""Exportação do inventário do Mapa dos contentores (PDF e Excel) para
imprimir — organizado por contentor → canister → andar.

Funções puras sobre o DataFrame de `obter_inventario_contentores`
(stock_repo): sem BD nem Streamlit, para poderem correr dentro do
callable de `st.download_button`.
"""

from __future__ import annotations

import io
import re
from xml.sax.saxutils import escape
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

from modules.mapa_layout import chave_natural

# (coluna do DataFrame, cabeçalho) — mesma ordem no PDF e no Excel.
COLUNAS_LOTE = [
    ("andar", "Andar"),
    ("garanhao", "Garanhão"),
    ("proprietario", "Proprietário"),
    ("palhetas", "Palhetas"),
    ("referencia", "Data / Ref."),
    ("cor", "Cor"),
    ("qualidade", "Qualidade"),
    ("motilidade", "Mot. %"),
    ("concentracao", "Conc. M/ml"),
    ("observacoes", "Observações"),
]

_CINZA_ESCURO = "#334155"
_CINZA_CLARO = "#f1f5f9"
_LINHA = "#cbd5e1"


def texto_celula(valor) -> str:
    """Valor de célula como texto; vazios (None, NaN, NaT, pd.NA dos
    inteiros `Int64`) dão ""."""
    if valor is None or (pd.api.types.is_scalar(valor) and pd.isna(valor)):
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def texto_pdf(valor) -> str:
    """Texto para um `Paragraph` do reportlab, que interpreta marcação
    XML — sem escapar, um "&" ou "<" num nome partia a geração do PDF."""
    return escape(texto_celula(valor))


def ordenar_inventario(df: pd.DataFrame) -> pd.DataFrame:
    """Ordena por contentor (ordem natural do código: 1, 2, …, 10, nomes),
    canister, andar e garanhão."""
    if df.empty:
        return df.copy()
    out = df.copy()
    out["_ordem_cont"] = out["contentor"].map(chave_natural)
    out = out.sort_values(["_ordem_cont", "canister", "andar", "garanhao"], na_position="last")
    return out.drop(columns="_ordem_cont").reset_index(drop=True)


def _contentores(df: pd.DataFrame):
    """[(codigo, linhas_com_lote)] pela ordem do inventário. Um contentor
    vazio vem com `linhas` vazio."""
    df = ordenar_inventario(df)
    grupos = []
    for codigo in dict.fromkeys(df["contentor"]):
        linhas = df[(df["contentor"] == codigo) & df["palhetas"].notna()]
        grupos.append((codigo, linhas))
    return grupos


def _total(linhas: pd.DataFrame) -> int:
    return int(pd.to_numeric(linhas["palhetas"], errors="coerce").fillna(0).sum())


# ------------------------------------------------------------------
# PDF
# ------------------------------------------------------------------

class _DocInventario(SimpleDocTemplate):
    """Escreve "Contentor X — continuação" no topo das páginas em que um
    contentor continua da página anterior (sem o título dele), para uma
    folha solta impressa nunca ficar sem saber de que contentor é."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._contentor_atual = None
        self._titulo_nesta_pagina = False

    def afterFlowable(self, flowable):
        codigo = getattr(flowable, "_contentor", None)
        if codigo is not None:
            self._contentor_atual = codigo
            self._titulo_nesta_pagina = True

    def afterPage(self):
        if self._contentor_atual is not None and not self._titulo_nesta_pagina:
            c = self.canv
            c.saveState()
            c.setFont("Helvetica", 8)
            c.setFillColor(colors.HexColor("#64748b"))
            c.drawString(self.leftMargin, self.pagesize[1] - 0.8 * cm,
                         f"Contentor {self._contentor_atual} — continuação")
            c.restoreState()
        self._titulo_nesta_pagina = False


def gerar_pdf_inventario(df: pd.DataFrame, gerado_em: datetime | None = None) -> bytes:
    """PDF A4 horizontal: uma página por contentor; dentro, um bloco por
    canister com a tabela dos lotes ordenada por andar."""
    gerado_em = gerado_em or datetime.now()
    buffer = io.BytesIO()
    doc = _DocInventario(
        buffer, pagesize=landscape(A4),
        leftMargin=1.4 * cm, rightMargin=1.4 * cm, topMargin=1.3 * cm, bottomMargin=1.3 * cm,
        title="Inventário dos contentores",
    )
    base = getSampleStyleSheet()
    st_titulo = ParagraphStyle("t", parent=base["Heading1"], fontSize=16, spaceAfter=2,
                               textColor=colors.HexColor(_CINZA_ESCURO))
    st_sub = ParagraphStyle("s", parent=base["Normal"], fontSize=9, textColor=colors.HexColor("#64748b"),
                            spaceAfter=10)
    st_canister = ParagraphStyle("c", parent=base["Heading3"], fontSize=11, spaceBefore=8, spaceAfter=4,
                                 textColor=colors.HexColor(_CINZA_ESCURO))
    st_celula = ParagraphStyle("cel", parent=base["Normal"], fontSize=8, leading=10)
    st_vazio = ParagraphStyle("v", parent=base["Normal"], fontSize=10, textColor=colors.HexColor("#64748b"))

    # Larguras (cm) na ordem de COLUNAS_LOTE — somam a largura útil (~26.9).
    larguras = [1.4, 4.2, 4.2, 1.8, 2.6, 1.8, 2.4, 1.5, 2.0, 5.0]
    numericas = {"andar", "palhetas", "motilidade", "concentracao"}

    historia = []
    grupos = _contentores(df)
    if not grupos:
        historia.append(Paragraph("Inventário dos contentores", st_titulo))
        historia.append(Paragraph("Sem contentores para exportar.", st_vazio))

    for i, (codigo, linhas) in enumerate(grupos):
        if i:
            historia.append(PageBreak())
        titulo = Paragraph(f"Contentor {texto_pdf(codigo)}", st_titulo)
        titulo._contentor = texto_celula(codigo)
        historia.append(titulo)
        historia.append(Paragraph(
            f"{len(linhas)} lote(s) · {_total(linhas)} palhetas · "
            f"gerado em {gerado_em:%d/%m/%Y %H:%M}", st_sub))

        if linhas.empty:
            historia.append(Paragraph("Sem lotes neste contentor.", st_vazio))
            continue

        for canister, lc in linhas.groupby("canister", sort=False, dropna=False):
            nome_can = f"Canister {texto_celula(canister)}" if texto_celula(canister) else "Sem canister"
            dados = [[h for _c, h in COLUNAS_LOTE]]
            for _, r in lc.iterrows():
                linha = []
                for col, _h in COLUNAS_LOTE:
                    v = texto_celula(r.get(col))
                    if col == "andar" and v:
                        v = f"{v}º"
                    linha.append(v if col in numericas else Paragraph(escape(v), st_celula))
                dados.append(linha)
            tabela = Table(dados, colWidths=[w * cm for w in larguras], repeatRows=1)
            estilo = [
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(_CINZA_CLARO)),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor(_CINZA_ESCURO)),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor(_LINHA)),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
            for j, (col, _h) in enumerate(COLUNAS_LOTE):
                if col in numericas:
                    estilo.append(("ALIGN", (j, 1), (j, -1), "RIGHT" if col != "andar" else "CENTER"))
            tabela.setStyle(TableStyle(estilo))
            cabecalho = Paragraph(f"{nome_can} — {_total(lc)} palhetas", st_canister)
            # Cabeçalho do canister nunca fica sozinho no fundo da página;
            # uma tabela maior do que a página parte-se na mesma.
            historia.append(KeepTogether([cabecalho, tabela]))
            historia.append(Spacer(1, 4))

    def _rodape(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#94a3b8"))
        canvas.drawRightString(doc_.pagesize[0] - doc_.rightMargin, 0.7 * cm, f"Página {doc_.page}")
        canvas.restoreState()

    doc.build(historia, onFirstPage=_rodape, onLaterPages=_rodape)
    return buffer.getvalue()


# ------------------------------------------------------------------
# Excel
# ------------------------------------------------------------------

_NOME_FOLHA_INVALIDO = re.compile(r"[\[\]\*\?/\\:]")


def _nome_folha(codigo, usados: set) -> str:
    base = _NOME_FOLHA_INVALIDO.sub("-", texto_celula(codigo) or "Sem código")[:31] or "Contentor"
    nome, n = base, 2
    while nome.lower() in usados:
        sufixo = f" ({n})"
        nome, n = base[: 31 - len(sufixo)] + sufixo, n + 1
    usados.add(nome.lower())
    return nome


def _escrever_folha(ws, linhas: pd.DataFrame, com_contentor: bool):
    colunas = ([("contentor", "Contentor")] if com_contentor else []) + [("canister", "Canister")] + COLUNAS_LOTE
    cab_fill = PatternFill("solid", fgColor=_CINZA_CLARO.lstrip("#"))
    borda = Border(bottom=Side(style="thin", color=_LINHA.lstrip("#")))
    ws.append([h for _c, h in colunas])
    for cel in ws[1]:
        cel.font = Font(bold=True, color=_CINZA_ESCURO.lstrip("#"))
        cel.fill = cab_fill
        cel.border = borda
        cel.alignment = Alignment(vertical="center")

    for _, r in linhas.iterrows():
        valores = []
        for col, _h in colunas:
            v = r.get(col)
            if v is None or (isinstance(v, float) and pd.isna(v)):
                valores.append(None)
            elif col in {"canister", "andar", "palhetas", "motilidade", "concentracao"}:
                valores.append(int(v))
            else:
                valores.append(texto_celula(v))
        ws.append(valores)

    if not linhas.empty:
        idx_palhetas = [c for c, _h in colunas].index("palhetas") + 1
        letra = get_column_letter(idx_palhetas)
        ultima = ws.max_row
        ws.append([])
        total = ["Total"] + [None] * (len(colunas) - 1)
        ws.append(total)
        ws.cell(row=ws.max_row, column=idx_palhetas, value=f"=SUM({letra}2:{letra}{ultima})")
        for cel in ws[ws.max_row]:
            cel.font = Font(bold=True)

    larguras = {"contentor": 14, "canister": 10, "andar": 8, "garanhao": 26, "proprietario": 26,
                "palhetas": 10, "referencia": 14, "cor": 10, "qualidade": 14, "motilidade": 9,
                "concentracao": 11, "observacoes": 40}
    for i, (col, _h) in enumerate(colunas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = larguras.get(col, 12)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{max(1, ws.max_row)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def gerar_excel_inventario(df: pd.DataFrame) -> bytes:
    """Excel com uma folha "Todos" (todos os contentores) e uma folha por
    contentor, ordenadas por canister e andar; prontas a imprimir
    (horizontal, a caber na largura, cabeçalho repetido)."""
    wb = Workbook()
    ws_todos = wb.active
    ws_todos.title = "Todos"
    grupos = _contentores(df)
    todas = pd.concat([l for _c, l in grupos]) if grupos else pd.DataFrame(columns=df.columns)
    _escrever_folha(ws_todos, todas, com_contentor=True)

    usados = {"todos"}
    for codigo, linhas in grupos:
        _escrever_folha(wb.create_sheet(_nome_folha(codigo, usados)), linhas, com_contentor=False)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
