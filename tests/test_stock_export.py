"""Testes para a exportação do "Stock completo" (Relatórios → Histórico
geral) em PDF, CSV e Excel."""

from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

from modules.stock_export import (
    COLUNAS_LOTES, gerar_csv_stock, gerar_excel_stock, gerar_pdf_stock, tabelas_stock,
)


def _stock():
    """Formato de `carregar_stock` (e.* + proprietario_nome,
    contentor_codigo, garanhao_nome)."""
    return pd.DataFrame([
        dict(id=1, garanhao="rubi", garanhao_nome="Rubi", proprietario_nome="Coudelaria A & B",
             data_embriovet="2021", origem_externa=None, palhetas_produzidas=40, quantidade_inicial=40,
             existencia_atual=12, qualidade="Boa", cor="Azul", concentracao=200, motilidade=60,
             certificado="Sim", dose="8", contentor_codigo="1", canister=3, andar=2,
             local_armazenagem=None, observacoes="<frágil>", criado_por="maria",
             data_criacao=datetime(2026, 10, 1, 9, 30)),
        dict(id=2, garanhao="xaquiro", garanhao_nome="Xaquiro", proprietario_nome=None,
             data_embriovet=None, origem_externa="EXT-1", palhetas_produzidas=None, quantidade_inicial=None,
             existencia_atual=5, qualidade=None, cor=None, concentracao=None, motilidade=None,
             certificado="Não", dose=None, contentor_codigo="TEMPORÁRIO", canister=1, andar=1,
             local_armazenagem=None, observacoes=None, criado_por=None, data_criacao=None),
        dict(id=3, garanhao="rubi", garanhao_nome="Rubi", proprietario_nome="João",
             data_embriovet="2022", origem_externa=None, palhetas_produzidas=10, quantidade_inicial=10,
             existencia_atual=3, qualidade="Média", cor="Verde", concentracao=150, motilidade=50,
             certificado="Sim", dose=None, contentor_codigo="1", canister=4, andar=1,
             local_armazenagem=None, observacoes=None, criado_por="andre", data_criacao=None),
    ])


def test_tabelas_lotes_tem_todos_os_campos():
    t = tabelas_stock(_stock())
    assert list(t["Lotes"].columns) == [h for _c, h in COLUNAS_LOTES]
    assert len(t["Lotes"]) == 3
    # Ordenado por garanhão e proprietário; data formatada.
    assert list(t["Lotes"]["Proprietário"].fillna("—")) == ["Coudelaria A & B", "João", "—"]
    assert t["Lotes"].iloc[0]["Data de criação"] == "01/10/2026 09:30"


def test_resumos():
    t = tabelas_stock(_stock())
    g = t["Por garanhão"].set_index("Garanhão")
    assert int(g.loc["Rubi", "Palhetas"]) == 15 and int(g.loc["Rubi", "Lotes"]) == 2
    assert int(g.loc["Rubi", "Proprietários"]) == 2
    p = t["Por proprietário"].set_index("Proprietário")
    assert int(p.loc["—", "Palhetas"]) == 5  # sem proprietário
    c = t["Por contentor"].set_index("Contentor")
    assert int(c.loc["1", "Canisters"]) == 2


def test_csv():
    texto = gerar_csv_stock(_stock()).decode("utf-8-sig")
    linhas = texto.strip().splitlines()
    assert linhas[0].startswith("Garanhão,Proprietário,Data produção")
    assert len(linhas) == 4


def test_excel():
    wb = load_workbook(io.BytesIO(gerar_excel_stock(_stock())))
    assert wb.sheetnames == ["Lotes", "Por garanhão", "Por proprietário", "Por contentor"]
    lotes = wb["Lotes"]
    assert [c.value for c in lotes[1]][:3] == ["Garanhão", "Proprietário", "Data produção"]
    assert lotes.cell(row=6, column=1).value == "Total"


def test_texto_celula_vazios():
    from modules.mapa_export import texto_celula
    assert [texto_celula(v) for v in (None, float("nan"), pd.NA, pd.NaT, 3.0, " x ")] == ["", "", "", "", "3", "x"]


def test_pdf_com_caracteres_especiais():
    pdf = gerar_pdf_stock(_stock(), "período 01/10/2026 a 06/10/2026")
    assert pdf.startswith(b"%PDF")


def test_vazio():
    vazio = _stock().iloc[0:0]
    assert gerar_pdf_stock(vazio).startswith(b"%PDF")
    assert load_workbook(io.BytesIO(gerar_excel_stock(vazio))).sheetnames[0] == "Lotes"
    assert gerar_csv_stock(vazio).decode("utf-8-sig").startswith("Garanhão")
