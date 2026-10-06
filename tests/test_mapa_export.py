"""Testes para a exportação do inventário do Mapa (PDF/Excel), organizado
por contentor → canister → andar."""

from __future__ import annotations

import io
import os
import re
import time

import pandas as pd
import psycopg2
from openpyxl import load_workbook

from modules.mapa_export import gerar_excel_inventario, gerar_pdf_inventario, ordenar_inventario


def _df():
    linhas = [
        # contentor, canister, andar, garanhão, proprietário, palhetas
        ("10", 2, 1, "Zeus", "Dono B", 4),
        ("2", 3, 2, "Atlas", "Dono A", 10),
        ("2", 1, 2, "Bravo", "Dono A", 5),
        ("2", 1, 1, "Corsario", None, 7),
        ("Equogestão", None, None, None, None, None),  # contentor vazio
    ]
    df = pd.DataFrame(linhas, columns=["contentor", "canister", "andar", "garanhao", "proprietario", "palhetas"])
    df["contentor_id"] = df["contentor"].map({"10": 10, "2": 2, "Equogestão": 99})
    for col in ["referencia", "cor", "qualidade", "motilidade", "concentracao", "observacoes"]:
        df[col] = None
    return df


def test_ordenar_contentor_canister_andar():
    out = ordenar_inventario(_df())
    chave = list(zip(out["contentor"], out["canister"], out["andar"]))
    assert chave[:4] == [("2", 1, 1), ("2", 1, 2), ("2", 3, 2), ("10", 2, 1)]
    assert out.iloc[-1]["contentor"] == "Equogestão"


def test_pdf_gerado():
    pdf = gerar_pdf_inventario(_df())
    assert pdf.startswith(b"%PDF")
    # Uma página por contentor (3); "/Type /Pages" é a raiz, não conta.
    assert len(re.findall(rb"/Type /Page\b(?!s)", pdf)) == 3


def test_pdf_sem_contentores():
    assert gerar_pdf_inventario(_df().iloc[0:0]).startswith(b"%PDF")


def test_excel_folhas_e_totais():
    wb = load_workbook(io.BytesIO(gerar_excel_inventario(_df())))
    assert wb.sheetnames == ["Todos", "2", "10", "Equogestão"]

    todos = wb["Todos"]
    cab = [c.value for c in todos[1]]
    assert cab[:4] == ["Contentor", "Canister", "Andar", "Garanhão"]
    # 4 lotes (o contentor vazio não gera linha de lote), linha em branco, total.
    assert [todos.cell(row=r, column=1).value for r in range(2, 6)] == ["2", "2", "2", "10"]
    assert todos.cell(row=7, column=1).value == "Total"
    col_palhetas = cab.index("Palhetas") + 1
    assert todos.cell(row=7, column=col_palhetas).value.startswith("=SUM(")

    folha2 = wb["2"]
    assert [folha2.cell(row=r, column=1).value for r in range(2, 5)] == [1, 1, 3]  # canister
    assert wb["Equogestão"].max_row == 1  # só cabeçalho


def test_obter_inventario_contentores_bd():
    from modules.repositories.stock_repo import obter_inventario_contentores

    conn = psycopg2.connect(os.getenv("DATABASE_URL", "").strip())
    cur = conn.cursor()
    sufixo = int(time.time() * 1_000_000)
    cur.execute(
        "INSERT INTO contentores (codigo, x, y, w, h, ativo) VALUES (%s, 0, 0, 70, 70, TRUE) RETURNING id",
        (f"_TEST_EXP_{sufixo}",),
    )
    cid = cur.fetchone()[0]
    cur.execute(
        "INSERT INTO contentores (codigo, x, y, w, h, ativo) VALUES (%s, 0, 0, 70, 70, TRUE) RETURNING id",
        (f"_TEST_EXP_VAZIO_{sufixo}",),
    )
    cid_vazio = cur.fetchone()[0]
    cur.execute(
        "INSERT INTO estoque_dono (garanhao, existencia_atual, contentor_id, canister, andar) "
        "VALUES (%s, 6, %s, 4, 2), (%s, 0, %s, 1, 1) RETURNING id",
        (f"_TEST_EXP_G_{sufixo}", cid, f"_TEST_EXP_ZERO_{sufixo}", cid),
    )
    lotes = [r[0] for r in cur.fetchall()]
    conn.commit()
    try:
        df = obter_inventario_contentores([cid, cid_vazio])
        assert set(df["contentor_id"]) == {cid, cid_vazio}
        com = df[df["contentor_id"] == cid]
        assert len(com) == 1  # o lote com existência 0 fica de fora
        assert (int(com.iloc[0]["canister"]), int(com.iloc[0]["andar"]), int(com.iloc[0]["palhetas"])) == (4, 2, 6)
        assert df[df["contentor_id"] == cid_vazio]["palhetas"].isna().all()
    finally:
        cur.execute("DELETE FROM estoque_dono WHERE id = ANY(%s)", (lotes,))
        cur.execute("DELETE FROM contentores WHERE id = ANY(%s)", ([cid, cid_vazio],))
        conn.commit()
        cur.close()
        conn.close()
